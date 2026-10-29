"""Spreadsheet and CSV lead import: analyse, map columns, commit, review rejected rows
and roll back a batch.
"""
import csv
import io
import os
import uuid
from datetime import date, datetime
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import Response
from pydantic import BaseModel

import crm_schema
import pipecat_server as core

router = APIRouter()


class ImportRow(BaseModel):
    name: str
    organization: Optional[str] = "Company"
    email: Optional[str] = ""
    mobile: Optional[str] = ""
    policyType: Optional[str] = "Commercial Fleet Insurance"
    premium: Optional[str] = "₹ 45,000"
    status: Optional[str] = "New"
    assignedTo: Optional[str] = ""
    expiryDate: Optional[str] = ""
    policyNumber: Optional[str] = ""


class ImportPayload(BaseModel):
    leads: List[ImportRow]
    deduplicationStrategy: Optional[str] = "skip"


@router.post("/api/leads/import")
def import_leads(req: ImportPayload):
    db = core.load_db()
    imported = 0
    updated = 0
    skipped = 0
    errors = []
    now_str = datetime.utcnow().isoformat() + "Z"

    for row_number, r in enumerate(req.leads, start=1):
        full_name = r.name.strip()
        if not full_name or not (r.mobile or r.email):
            errors.append({"row":row_number,"message":"Name and either mobile or email are required"}); continue
        org_name = r.organization.strip() if r.organization else "Corporate Client"
        
        existing = next((l for l in db.get("leads", []) if l.get("mobile") == r.mobile or (r.email and l.get("email") == r.email)), None)
        if existing and req.deduplicationStrategy == "skip":
            skipped += 1
            continue
        elif existing and req.deduplicationStrategy == "update":
            existing.update({
                "organization": org_name,
                "policyType": r.policyType or existing.get("policyType"),
                "premium": r.premium or existing.get("premium"),
                "lastModified": now_str
            })
            # Re-importing a refreshed book of business must move the renewal date forward,
            # otherwise the radar keeps tracking last year's expiry.
            if r.expiryDate:
                existing["expiryDate"] = r.expiryDate
                existing["daysRemaining"] = core._days_until(r.expiryDate)
            if r.policyNumber:
                existing["policyNumber"] = r.policyNumber
            updated += 1
            continue

        lid = f"L{len(db.get('leads', [])) + 1}_{uuid.uuid4().hex[:4]}"
        new_lead = {
            "id": lid,
            "salutation": "Mr",
            "firstName": full_name.split(" ")[0],
            "lastName": " ".join(full_name.split(" ")[1:]) if " " in full_name else "",
            "name": full_name,
            "organization": org_name,
            "website": "",
            "status": r.status or "New",
            "email": r.email or "",
            "mobile": r.mobile,
            "gender": "Other",
            "territory": "West India",
            "industry": "Commercial",
            "annualRevenue": "₹ 1,00,00,000",
            "noOfEmployees": "11-50",
            "companyDescription": "Imported via Spreadsheet batch",
            "policyType": r.policyType or "Commercial Fleet Insurance",
            "policyNumber": r.policyNumber or f"POL-{uuid.uuid4().hex[:6].upper()}",
            "premium": r.premium or "₹ 45,000",
            "expiryDate": r.expiryDate or "",
            "daysRemaining": core._days_until(r.expiryDate) if r.expiryDate else None,
            "nextAction": "Initial Discovery Discussion",
            "nextActionDate": datetime.now().strftime("%Y-%m-%d"),
            "nextActionCompleted": False,
            "source": "Spreadsheet Import",
            "assignedTo": r.assignedTo or core.get_next_round_robin_user(db),
            "lastModified": now_str,
            "lostReason": ""
        }
        db.setdefault("leads", []).insert(0, new_lead)
        imported += 1

    core.audit(db,"Lead Import Completed","Import",f"IMP_{uuid.uuid4().hex[:8]}",{"imported":imported,"updated":updated,"skipped":skipped,"errors":errors})
    core.save_db(db)
    return {"status": "success", "importedCount": imported, "updatedCount": updated, "skippedCount":skipped, "errors":errors}


class LeadContactIndex:
    """Phone/email lookup for bulk work.

    Scanning every lead per imported row is quadratic: a 2,000-row file against a
    5,000-lead book is 10 million comparisons. Build the index once instead, and keep
    it current as rows are added.
    """

    def __init__(self, leads: List[Dict[str, Any]]):
        self.by_phone: Dict[str, Dict[str, Any]] = {}
        self.by_email: Dict[str, Dict[str, Any]] = {}
        for lead in leads:
            self.add(lead)

    def add(self, lead: Dict[str, Any]):
        phone = core._phone_key(lead.get("mobile"))
        if phone:
            self.by_phone.setdefault(phone, lead)
        email = (lead.get("email") or "").strip().lower()
        if email:
            self.by_email.setdefault(email, lead)

    def find(self, mobile: str = "", email: str = "") -> Optional[Dict[str, Any]]:
        hit = self.by_phone.get(core._phone_key(mobile)) if mobile else None
        if hit:
            return hit
        return self.by_email.get((email or "").strip().lower()) if email else None


MAX_IMPORT_ROWS = 5000


MAX_IMPORT_BYTES = 10 * 1024 * 1024


def _normalise_header(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").replace(".", " ").split())


def _suggest_field(header: str, taken: set, custom_labels: Optional[Dict[str, str]] = None) -> Optional[str]:
    """Match a spreadsheet column to a CRM field: exact alias first, then a custom field's
    exact label, then containment."""
    clean = _normalise_header(header)
    if not clean:
        return None
    for spec in core.IMPORT_FIELDS:
        if spec["field"] in taken:
            continue
        if clean in spec["aliases"] or clean == spec["field"].lower():
            return spec["field"]
    for label, target in (custom_labels or {}).items():
        if clean == label and target not in taken:
            return target
    for spec in core.IMPORT_FIELDS:
        if spec["field"] in taken:
            continue
        for alias in spec["aliases"]:
            if len(alias) >= 4 and (alias in clean or clean in alias):
                return spec["field"]
    return None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_spreadsheet(content: bytes, filename: str) -> Dict[str, Any]:
    """Return {headers, rows} from CSV/TSV or XLSX bytes."""
    extension = os.path.splitext(filename)[1].lower()
    grid: List[List[Any]] = []

    if extension in [".xlsx", ".xlsm"]:
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise HTTPException(status_code=422, detail="Excel support is unavailable on this server. Save the file as CSV and retry.")
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(status_code=422, detail="That file could not be read as a spreadsheet. Check it is a valid .xlsx file.")
        sheet = workbook[workbook.sheetnames[0]]
        for row in sheet.iter_rows(values_only=True):
            grid.append(list(row))
            if len(grid) > MAX_IMPORT_ROWS + 50:
                break
        workbook.close()
    elif extension in [".csv", ".txt", ".tsv"]:
        for encoding in ["utf-8-sig", "utf-8", "latin-1"]:
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise HTTPException(status_code=422, detail="That file's text encoding could not be read. Re-save it as UTF-8 CSV.")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        grid = [row for row in csv.reader(io.StringIO(text), dialect)]
    else:
        raise HTTPException(status_code=422, detail="Upload a .csv, .tsv or .xlsx file.")

    grid = [row for row in grid if any(_cell_text(cell) for cell in row)]
    if not grid:
        raise HTTPException(status_code=422, detail="That spreadsheet has no readable rows.")
    if len(grid) - 1 > MAX_IMPORT_ROWS:
        raise HTTPException(status_code=413, detail=f"That file has more than {MAX_IMPORT_ROWS} rows. Split it and import in batches.")

    headers = [_cell_text(cell) for cell in grid[0]]
    width = max(len(row) for row in grid)
    headers += [""] * (width - len(headers))
    headers = [h or f"Column {i + 1}" for i, h in enumerate(headers)]
    rows = [[_cell_text(cell) for cell in row] + [""] * (width - len(row)) for row in grid[1:]]
    return {"headers": headers, "rows": rows}


@router.post("/api/imports/analyze")
async def analyze_import(request: Request, filename: str = "upload.csv"):
    """Step 1: read the file, propose a column mapping, and stage it for review."""
    content = await request.body()
    if not content:
        raise HTTPException(status_code=422, detail="The uploaded file is empty.")
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="That file is larger than 10 MB.")

    parsed = _parse_spreadsheet(content, filename)
    headers, rows = parsed["headers"], parsed["rows"]

    db = core.load_db()
    custom_fields = [d for d in core.custom_field_defs(db) if not d.get("hidden")]
    custom_labels = {_normalise_header(d["label"]): f"custom.{d['key']}" for d in custom_fields}
    taken: set = set()
    columns = []
    for index, header in enumerate(headers):
        suggestion = _suggest_field(header, taken, custom_labels)
        if suggestion:
            taken.add(suggestion)
        samples = [r[index] for r in rows[:5] if index < len(r) and r[index]][:3]
        filled = sum(1 for r in rows if index < len(r) and r[index])
        columns.append({"index": index, "header": header, "samples": samples,
                        "filledCount": filled, "suggestedField": suggestion})

    batch = {
        "id": f"IMP_{uuid.uuid4().hex[:8]}", "filename": os.path.basename(filename),
        "status": "Awaiting mapping", "rowCount": len(rows), "columns": columns,
        "headers": headers, "rows": rows, "createdAt": datetime.utcnow().isoformat() + "Z",
        "importedCount": 0, "updatedCount": 0, "skippedCount": 0, "errors": [], "createdLeadIds": []
    }
    db.setdefault("importBatches", []).insert(0, batch)
    # Staged files are working data, not history: keep the most recent few only.
    db["importBatches"] = db["importBatches"][:25]
    core.audit(db, "Import File Analysed", "Import", batch["id"], {"filename": batch["filename"], "rows": len(rows)})
    core.save_db(db)

    return {"importId": batch["id"], "filename": batch["filename"], "rowCount": len(rows),
            "columns": columns, "previewRows": rows[:8],
            "availableFields": [{"field": f["field"], "label": f["label"], "required": f["required"]} for f in core.IMPORT_FIELDS]
                               + [{"field": f"custom.{d['key']}", "label": d["label"], "required": bool(d.get("required")),
                                   "custom": True} for d in custom_fields]}


class ImportCommit(BaseModel):
    mapping: Dict[str, str]                      # column index (as string) → CRM field, "__custom__" or "__skip__"
    deduplicationStrategy: Optional[str] = "skip"
    defaults: Optional[Dict[str, str]] = None


@router.post("/api/imports/{import_id}/commit")
def commit_import(import_id: str, payload: ImportCommit):
    """Step 2: apply the confirmed mapping. Unmapped columns are retained as custom fields."""
    db = core.load_db()
    batch = core._find_record(db.get("importBatches", []), import_id, "Import batch")
    if batch.get("status") == "Imported":
        raise HTTPException(status_code=409, detail="This file has already been imported.")

    headers, rows = batch.get("headers", []), batch.get("rows", [])
    mapping = {int(k): v for k, v in payload.mapping.items() if v and v != "__skip__"}
    if "name" not in mapping.values():
        raise HTTPException(status_code=422, detail="Map a column to Client name before importing.")
    custom_defs = core.custom_field_defs(db)
    unknown = [t for t in mapping.values() if t.startswith("custom.") and t[len("custom."):] not in {d["key"] for d in custom_defs}]
    if unknown:
        raise HTTPException(status_code=422, detail=f"Not a custom field: {', '.join(sorted(set(unknown)))}")

    defaults = payload.defaults or {}
    strategy = payload.deduplicationStrategy or "skip"
    now_str = datetime.utcnow().isoformat() + "Z"
    imported = updated = skipped = 0
    errors: List[Dict[str, Any]] = []
    created_ids: List[str] = []
    contact_index = LeadContactIndex(db.get("leads", []))     # one pass, not one per row
    org_names = {o.get("name", "").lower() for o in db.get("organizations", [])}
    contact_phones = {core._phone_key(c.get("phone")) for c in db.get("contacts", [])}

    for row_number, row in enumerate(rows, start=2):   # row 1 is the header
        values: Dict[str, str] = {}
        custom: Dict[str, str] = {}                # unmapped columns, kept as they were typed
        typed: Dict[str, str] = {}                 # columns mapped to the company's own fields
        for index, target in mapping.items():
            cell = row[index] if index < len(row) else ""
            if target == "__custom__":
                if cell:
                    custom[headers[index] if index < len(headers) else f"Column {index+1}"] = cell
            elif target.startswith("custom."):
                if cell:
                    typed[target[len("custom."):]] = cell
            elif cell:
                values[target] = cell

        name = (values.get("name") or "").strip()
        mobile = (values.get("mobile") or "").strip()
        email = (values.get("email") or "").strip()
        if not name:
            errors.append({"row": row_number, "message": "Client name is empty", "data": dict(zip(headers, row))})
            continue
        if not mobile and not email:
            errors.append({"row": row_number, "message": "Row has neither a mobile number nor an email address", "data": dict(zip(headers, row))})
            continue

        expiry = core._parse_flexible_date(values.get("expiryDate"))
        if values.get("expiryDate") and not expiry:
            errors.append({"row": row_number, "message": f"Could not read the renewal date '{values.get('expiryDate')}'", "data": dict(zip(headers, row))})
            continue

        existing = None
        if strategy != "create":
            existing = contact_index.find(mobile, email)
        if existing and strategy == "skip":
            skipped += 1
            continue
        if existing and strategy == "update":
            merged, problems = crm_schema.clean_custom(custom_defs, typed, existing.get("custom"),
                                                       product=existing.get("product") or existing.get("policyType"))
            if problems:
                errors.append({"row": row_number, "message": "; ".join(problems.values()), "data": dict(zip(headers, row))})
                continue
            if typed:
                existing["custom"] = merged
            for field, value in values.items():
                if field != "name" and value:
                    existing[field] = value
            if expiry:
                existing["expiryDate"] = expiry
                existing["daysRemaining"] = core._days_until(expiry)
            if custom:
                existing.setdefault("customFields", {}).update(custom)
            existing["lastModified"] = now_str
            existing["importBatchId"] = import_id
            updated += 1
            continue

        typed_custom, problems = crm_schema.clean_custom(custom_defs, typed, None, creating=True,
                                                         product=values.get("policyType") or defaults.get("policyType"))
        if problems:
            errors.append({"row": row_number, "message": "; ".join(problems.values()), "data": dict(zip(headers, row))})
            continue

        stage = values.get("status") if values.get("status") in core.STAGES else (defaults.get("status") or "New")
        lead_id = f"L_{uuid.uuid4().hex[:8]}"
        owner = values.get("assignedTo") or defaults.get("assignedTo") or core.apply_assignment_rules(db, {
            "id": lead_id, "name": name,
            "policyType": values.get("policyType") or defaults.get("policyType", ""),
            "source": f"Spreadsheet import ({batch.get('filename')})",
            "organization": values.get("organization") or defaults.get("organization", ""),
            "premium": values.get("premium", ""), "status": stage,
            "territory": values.get("territory") or defaults.get("territory", ""),
            "industry": values.get("industry") or defaults.get("industry", "")}) or core.get_next_round_robin_user(db)
        parts = name.split(" ")
        new_lead = {
            "id": lead_id, "salutation": "", "firstName": parts[0], "lastName": " ".join(parts[1:]),
            "name": name, "organization": values.get("organization") or defaults.get("organization") or "",
            "website": "", "status": stage, "email": email, "mobile": mobile, "gender": "",
            "territory": values.get("territory") or defaults.get("territory") or "",
            "industry": values.get("industry") or defaults.get("industry") or "",
            "annualRevenue": "", "noOfEmployees": "",
            "companyDescription": values.get("companyDescription") or "",
            "policyType": values.get("policyType") or defaults.get("policyType") or "",
            "policyNumber": values.get("policyNumber") or "",
            "insurer": values.get("insurer") or "",
            "premium": values.get("premium") or "", "sumInsured": values.get("sumInsured") or "",
            "startDate": core._parse_flexible_date(values.get("startDate")) or "",
            "expiryDate": expiry or "", "daysRemaining": core._days_until(expiry) if expiry else None,
            "nextAction": "Review imported account", "nextActionDate": datetime.now().strftime("%Y-%m-%d"),
            "nextActionCompleted": False, "source": f"Spreadsheet import ({batch.get('filename')})",
            "assignedTo": owner, "lastModified": now_str, "lostReason": "",
            "customFields": custom, "importBatchId": import_id, "importedRow": row_number,
            **({"custom": typed_custom} if typed_custom else {}),
            "assignedByRule": core._last_rule_match["name"]
        }
        task_id = f"T_{uuid.uuid4().hex[:8]}"
        new_lead["nextActionTaskId"] = task_id
        new_lead["createdAt"] = now_str
        new_lead["initialName"] = name
        db.setdefault("tasks", []).insert(0, {
            "id": task_id, "title": new_lead["nextAction"], "associatedType": "Lead",
            "associatedId": lead_id, "associatedTitle": f"{new_lead['name']} ({new_lead['organization'] or 'Individual'})",
            "priority": "Medium", "dueDate": new_lead["nextActionDate"], "status": "Pending",
            "assignedTo": owner, "activityType": "Follow-up", "createdAt": now_str, "updatedAt": now_str})
        db.setdefault("leads", []).insert(0, new_lead)
        contact_index.add(new_lead)          # so duplicates inside the same file are caught
        created_ids.append(lead_id)
        imported += 1

        org_name = new_lead["organization"]
        if org_name and org_name.lower() not in org_names:
            org_names.add(org_name.lower())
            db.setdefault("organizations", []).insert(0, {
                "id": f"O_{uuid.uuid4().hex[:8]}", "name": org_name, "website": "",
                "industry": new_lead["industry"], "annualRevenue": "", "territory": new_lead["territory"],
                "noOfEmployees": "", "lastModified": now_str})
        if mobile and core._phone_key(mobile) not in contact_phones:
            contact_phones.add(core._phone_key(mobile))
            db.setdefault("contacts", []).insert(0, {
                "id": f"C_{uuid.uuid4().hex[:8]}", "leadId": lead_id, "name": name, "email": email,
                "phone": mobile, "organization": org_name, "designation": "Primary Contact",
                "status": "Active", "lastModified": now_str})

    batch.update({"status": "Imported", "mapping": payload.mapping, "importedCount": imported,
                  "updatedCount": updated, "skippedCount": skipped, "errors": errors,
                  "createdLeadIds": created_ids, "committedAt": now_str,
                  "customFieldColumns": sorted({headers[i] for i, t in mapping.items() if t == "__custom__" and i < len(headers)})})
    batch.pop("rows", None)      # the staged copy is no longer needed once committed
    core.audit(db, "Import File Committed", "Import", import_id,
          {"imported": imported, "updated": updated, "skipped": skipped, "errors": len(errors)})
    core.save_db(db)

    return {"status": "success", "importedCount": imported, "updatedCount": updated,
            "skippedCount": skipped, "errors": errors, "createdLeadIds": created_ids}


@router.get("/api/imports")
def list_import_batches():
    batches = core.load_db().get("importBatches", [])
    return [{"id": b["id"], "filename": b.get("filename"), "status": b.get("status"),
             "rowCount": b.get("rowCount", 0), "importedCount": b.get("importedCount", 0),
             "updatedCount": b.get("updatedCount", 0), "skippedCount": b.get("skippedCount", 0),
             "createdAt": b.get("createdAt"), "committedAt": b.get("committedAt")} for b in batches]


@router.get("/api/imports/{import_id}")
def get_import_batch(import_id: str):
    batch = core._find_record(core.load_db().get("importBatches", []), import_id, "Import batch")
    return {k: v for k, v in batch.items() if k != "rows"}


@router.get("/api/imports/{import_id}/rejected.csv")
def download_rejected_rows(import_id: str):
    batch = core._find_record(core.load_db().get("importBatches", []), import_id, "Import batch")
    errors = batch.get("errors") or []
    if not errors:
        raise HTTPException(status_code=404, detail="This import had no rejected rows.")
    buffer = io.StringIO()
    columns = [c.get("header") for c in batch.get("columns", [])]
    writer = csv.writer(buffer)
    writer.writerow(["Row", "Reason"] + columns)
    for err in errors:
        writer.writerow([err.get("row"), err.get("message")] + [err.get("data", {}).get(c, "") for c in columns])
    return Response(content=buffer.getvalue(), media_type="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="rejected-{import_id}.csv"'})


@router.post("/api/imports/{import_id}/rollback")
def rollback_import(import_id: str):
    """Undo an import: archive the leads it created, leaving anything since worked on untouched."""
    db = core.load_db()
    batch = core._find_record(db.get("importBatches", []), import_id, "Import batch")
    if batch.get("rolledBackAt"):
        raise HTTPException(status_code=409, detail="This import has already been rolled back.")
    if batch.get("status") != "Imported":
        raise HTTPException(status_code=422, detail="Only a committed import can be rolled back.")

    created = set(batch.get("createdLeadIds", []))
    removed, kept = [], []
    for lead in list(db.get("leads", [])):
        if lead.get("id") not in created:
            continue
        is_worked = (
            lead.get("convertedPolicyId") or
            core.canonical_status(lead.get("status")) in core.CLOSED_STATUSES or
            (lead.get("initialName") and lead.get("name") != lead.get("initialName")) or
            (lead.get("createdAt") and lead.get("lastModified") and lead.get("lastModified") > lead.get("createdAt")) or
            any(n.get("leadId") == lead.get("id") or n.get("associatedId") == lead.get("id") for n in db.get("notes", []))
        )
        if is_worked:
            kept.append({"id": lead.get("id"), "name": lead.get("name"), "reason": "Already worked — left in place"})
            continue
        db["leads"] = [l for l in db["leads"] if l.get("id") != lead.get("id")]
        lead.update({"archived": True, "archivedReason": f"Rolled back import {import_id}", "archivedAt": datetime.utcnow().isoformat() + "Z"})
        db.setdefault("archivedLeads", []).insert(0, lead)
        removed.append(lead.get("id"))

    batch["rolledBackAt"] = datetime.utcnow().isoformat() + "Z"
    batch["status"] = "Rolled back"
    core.audit(db, "Import Rolled Back", "Import", import_id, {"archived": len(removed), "retained": len(kept)})
    core.save_db(db)
    return {"archivedCount": len(removed), "retained": kept}
